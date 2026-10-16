# -*- coding: utf-8 -*-
"""Fill the official Final Mission Package template for team revenge.
Headings, sub-headings and their order are taken verbatim from the template;
content is inserted beneath each one.

  python scripts/build-final-mission-package.py "<แบบฟอร์ม_Final_Mission_Package.docx>" \
         gistda-request/revenge_Final_Mission_Package.docx gistda-request/final-package-img 12
(needs: pip install python-docx · body font size in pt, TH Sarabun New — 12 pt gives 7 A4 pages)"""
import sys, copy, re
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

sys.stdout.reconfigure(encoding='utf-8')
TEMPLATE, OUT, IMG = sys.argv[1], sys.argv[2], sys.argv[3]
BODY = float(sys.argv[4]) if len(sys.argv) > 4 else 13
SMALL = BODY - 1
FONT = 'TH Sarabun New'
NAVY = '123F57'
TEXT = RGBColor(0x11, 0x11, 0x11)

doc = docx.Document(TEMPLATE)
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)       # A4 (template ships as Letter)
CONTENT_CM = 21.0 - 2 * 1.6

# ── low-level helpers ────────────────────────────────────────────────────────

def style_run(r, size=BODY, bold=False, color=TEXT, italic=False):
    r.font.name = FONT
    rpr = r._r.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rpr.insert(0, fonts)
    for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        fonts.set(qn(k), FONT)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    if color is not None:
        r.font.color.rgb = color
    # complex-script (Thai) size/bold must sit right after their Latin twins (schema order)
    szcs = rpr.find(qn('w:szCs'))
    if szcs is None:
        szcs = OxmlElement('w:szCs'); rpr.find(qn('w:sz')).addnext(szcs)
    szcs.set(qn('w:val'), str(int(size * 2)))
    if bold and rpr.find(qn('w:bCs')) is None:
        rpr.find(qn('w:b')).addnext(OxmlElement('w:bCs'))

def runs(p, text, size=BODY, color=TEXT, bold_all=False):
    """Add text to paragraph p; **…** marks bold spans."""
    for i, part in enumerate(re.split(r'\*\*', text)):
        if part:
            style_run(p.add_run(part), size=size, bold=bold_all or i % 2 == 1, color=color)

def fmt(p, indent=0.75, first=None, align='justify', after=1, size=BODY, keep=False):
    pf = p.paragraph_format
    pf.left_indent = Cm(indent)
    if first is not None:
        pf.first_line_indent = Cm(first)
    pf.space_before = Pt(0)
    pf.space_after = Pt(after)
    pf.line_spacing = Pt(size + 3)
    if keep:
        pf.keep_with_next = True
    # Thai has few spaces, so justified lines open wide gaps — body text stays left-aligned
    p.alignment = {'justify': WD_ALIGN_PARAGRAPH.LEFT, 'left': WD_ALIGN_PARAGRAPH.LEFT,
                   'center': WD_ALIGN_PARAGRAPH.CENTER}[align]

class Cursor:
    def __init__(self, anchor):
        self.el = anchor._p if hasattr(anchor, '_p') else anchor
    def put(self, el):
        self.el.addnext(el); self.el = el
        return el
    def para(self, text='', indent=0.75, first=None, align='justify', size=BODY, after=1, color=TEXT, bold=False, keep=False):
        p = Paragraph(OxmlElement('w:p'), doc._body)
        self.put(p._p)
        fmt(p, indent, first, align, after, size, keep)
        if text:
            runs(p, text, size=size, color=color, bold_all=bold)
        return p
    def bullets(self, items, indent=1.25, size=BODY, mark='•'):
        for it in items:
            p = self.para('', indent=indent, first=-0.4, size=size)
            style_run(p.add_run(mark + ' '), size=size)
            runs(p, it, size=size)
    def table(self, rows, widths_cm, header=True, size=SMALL, align=None, shade_first_col=False):
        t = doc.add_table(rows=len(rows), cols=len(widths_cm))
        t.style = doc.styles['Table Grid']
        self.put(t._tbl)
        layout(t, widths_cm)
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                fill_cell(cell, val, size=size, bold=header and i == 0, white=header and i == 0,
                          align=(align[j] if align else 'left') if not (header and i == 0) else 'center')
                if header and i == 0:
                    shade(cell, NAVY)
                elif shade_first_col and j == 0:
                    shade(cell, 'E8EEF4')
        if header:
            trpr = t.rows[0]._tr.get_or_add_trPr()
            h = OxmlElement('w:tblHeader'); trpr.append(h)
        # keep rows unbroken across pages
        for r in t.rows:
            trpr = r._tr.get_or_add_trPr()
            cs = OxmlElement('w:cantSplit'); trpr.append(cs)
        gap = self.para('', after=0, size=4)
        return t

def layout(t, widths_cm):
    tblPr = t._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    mar = OxmlElement('w:tblCellMar')
    for side, v in (('top', 10), ('left', 70), ('bottom', 10), ('right', 70)):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); mar.append(e)
    tblPr.append(mar)
    tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(int(sum(widths_cm) * 567))); tw.set(qn('w:type'), 'dxa')
    old = tblPr.find(qn('w:tblW'))
    if old is not None: tblPr.remove(old)
    tblPr.append(tw)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn('w:gridCol')), widths_cm):
        gc.set(qn('w:w'), str(int(w * 567)))
    for row in t.rows:
        for c, w in zip(row.cells, widths_cm):
            c.width = Cm(w)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:shd')): tcPr.remove(old)
    s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill)
    tcPr.append(s)

def fill_cell(cell, val, size=SMALL, bold=False, white=False, align='left'):
    """val: str (lines split on '|'), or ('img', path, width_cm, caption)."""
    for p in cell.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    p0 = cell.paragraphs[0]
    for r in list(p0.runs): r._r.getparent().remove(r._r)
    color = RGBColor(0xFF, 0xFF, 0xFF) if white else TEXT
    if isinstance(val, tuple) and val[0] == 'img':
        _, path, w, cap = val
        fmt(p0, 0, 0, 'center', 0, size)
        p0.paragraph_format.line_spacing = None
        p0.add_run().add_picture(path, width=Cm(w))
        pc = cell.add_paragraph(); fmt(pc, 0, 0, 'left', 0, size - 1)
        runs(pc, cap, size=size - 1)
        return
    lines = str(val).split('|')
    for k, line in enumerate(lines):
        p = p0 if k == 0 else cell.add_paragraph()
        fmt(p, 0, 0, align, 0, size)
        runs(p, line.strip(), size=size, color=color, bold_all=bold)

def find(prefix, start=0):
    ps = doc.paragraphs
    for i in range(start, len(ps)):
        if ps[i].text.strip().replace('​', '').startswith(prefix):
            return ps[i]
    raise KeyError(prefix)

def remove(p):
    p._p.getparent().remove(p._p)

# template headings stay glued to the content that follows them
HEAD = re.compile(r'^(\d+\.\s|\d+\.\d|\(อธิบาย|ระบุ|วางลิงก์|จุดเด่น)')
for p in doc.paragraphs:
    if HEAD.match(p.text.strip()):
        p.paragraph_format.keep_with_next = True

# ── cover: team table ───────────────────────────────────────────────────────
cover, dt, lm = doc.tables[0], doc.tables[1], doc.tables[2]   # grab before new tables shift the indices
COVER = [
    'revenge',
    'THEOS-2 Mae Sai Command : ระบบภารกิจจำลองสถานการณ์และสนับสนุนการตัดสินใจบริหารจัดการอุทกภัยแม่สาย ด้วยข้อมูลภูมิสารสนเทศแบบหลายช่วงคลื่น (Multi-Sensor Geospatial Serious Game and Integrated Disaster Decision Support System)',
    'โรงเรียนสิงห์บุรี',
    'นายสิทธิพัฒน์ จิตร์กระจ่าง (ผู้แทนทีม: นายศิรวิทย์ พิทักษ์สาลี และสมาชิก 5 คน)',
]
for row, val in zip(cover.rows, COVER):
    fill_cell(row.cells[1], val, size=14, bold=(val == 'revenge'))

# ── highlights ──────────────────────────────────────────────────────────────
hl = find('จุดเด่น/สำคัญของ Mission')
for p in list(doc.paragraphs):
    if p.text.strip().startswith('....') and len(p.text.strip()) > 50:
        remove(p)
c = Cursor(hl)
c.para('**THEOS-2 Mae Sai Command** คือ Serious Game บนเว็บที่ให้ผู้เรียนสวมบท “ผู้บัญชาการศูนย์ปฏิบัติการฉุกเฉิน (EOC) อำเภอแม่สาย” '
       'ตัดสินใจภายใต้หลักฐานที่ไม่เคยครบ ผ่าน 3 องก์ตามเหตุการณ์จริง ก.ย. 2567 — **องก์ 1 (T−48 ชม.)** เฝ้าระวังต้นน้ำข้ามพรมแดนด้วย THEOS-2 '
       'หารอยดินถล่ม วัดความกว้างลำน้ำ และอุดจุดฟันหลอของแนวตลิ่ง **องก์ 2 (T+0)** เมฆบัง 100% ต้องสลับไปใช้เรดาร์ SAR และ Sphere Routing API '
       'หาเส้นทางอพยพเลี่ยงน้ำท่วม **องก์ 3 (T+36 ชม.)** เยียวยารายหลังคาเรือนตามมติ ครม. (9,000 + 10,000 บาท) บนฐานข้อมูลอาคารจริงของ GISTDA 15,851 หลัง',
       indent=0, first=1.0, after=2)
c.para('ทุกการตัดสินใจขับเคลื่อนด้วยดัชนี **SFVI** ที่คำนวณสดจากข้อมูลหลายเซนเซอร์ ทุกตัวเลขมีป้ายบอกที่มา 4 สี และมีโหมดเรียนรู้ 2 ระดับ '
       '(**Beginner / Expert Mode**) ระบบบันทึกพฤติกรรมการใช้ข้อมูลของผู้เรียนเองเพื่อวัดผล Pre/Post-test ส่งออก CSV ได้ ทำงานได้ทั้งออนไลน์และออฟไลน์',
       indent=0, first=1.0, after=2)

# ── 1 · problem, area, target ───────────────────────────────────────────────
c = Cursor(find('1.1 ปัญหา'))
c.bullets([
    '**ขนาดและความเร็วของภัย:** อ.แม่สายน้ำท่วมซ้ำ 6 ครั้งในปี 2567 ข้อมูลดาวเทียม GISTDA ชี้พื้นที่ท่วมขังพุ่งจาก 6,182 ไร่ (10 ก.ย.) เป็น 25,204 ไร่ (12 ก.ย.) — เพิ่ม 4 เท่าใน 2 วัน เวลาเตือนภัยเหลือเพียง 2–4 ชม. และเมื่อน้ำลด โคลนทับถมทำให้อยู่อาศัยไม่ได้',
    '**สาเหตุอยู่ในที่ที่มองไม่เห็น:** ป่าต้นน้ำเปลี่ยนเป็นเหมืองแร่และเกษตร เกิดรอยดินถล่มกว้าง 20–30 ม. หลายจุด ขณะที่ลำน้ำสายช่วงปลายน้ำแคบลงจากราว 150 ม. เหลือ 20–30 ม.',
    '**ดาวเทียมเป็นทางเลือกเดียว:** ต้นน้ำแม่น้ำสายอยู่ในเมียนมา ไทยตั้งสถานีภาคพื้นไม่ได้ และ UAV บินข้ามพรมแดนไม่ได้',
    '**ช่องว่างการพิสูจน์สิทธิ์เยียวยา:** มติ ครม. 17 ก.ย. 67 กำหนดเหมาจ่าย 9,000 บาท/ครัวเรือน และค่าล้างโคลน 10,000 บาท/หลัง แต่การพิสูจน์รายบ้านทำได้ยาก ทั้งที่ดาวเทียมประเมินวันท่วมขังและโคลนได้ ผลสำรวจนักเรียนโรงเรียนสิงห์บุรี 80 คน พบว่า 20% ไม่ทราบว่าต้นน้ำอยู่นอกประเทศ และ 15% ไม่ทราบเกณฑ์เยียวยาตามวันท่วมขัง',
])
c = Cursor(find('1.2 พื้นที่ศึกษา'))
c.para('**AOI-A (เขตเมือง):** 4 ชุมชนเฝ้าระวังริมแม่น้ำสาย อ.แม่สาย จ.เชียงราย — สายลมจอย ม.1 ต.เวียงพางคำ, เกาะทราย ม.7 ต.แม่สาย, ไม้ลุงขน และเหมืองแดง '
       '**AOI-B (ต้นน้ำ):** พื้นที่รับน้ำต้นแม่น้ำสายในรัฐฉาน เมียนมา ราว 965 ตร.กม. เหตุผลที่เลือก: เป็นกรณีข้ามพรมแดนที่ตอบคำถาม “ทำไมต้องใช้ดาวเทียม” ชัดที่สุด '
       'มีข้อมูลอ้างอิงจริงจาก RADARSAT-2 (10, 12 ก.ย. 67) และ THEOS-2 (17 ก.ย. 67) และภัยเกิดซ้ำต่อเนื่อง', indent=1.25)
c = Cursor(find('1.3 กลุ่มเป้าหมาย'))
c.para('(1) นักเรียนมัธยมปลาย นักศึกษา และครู (2) ผู้เข้าชมแหล่งเรียนรู้/นิทรรศการ Roadshow (3) เจ้าหน้าที่ท้องถิ่นและชุมชนในพื้นที่เสี่ยงภัย สำหรับซ้อมแผนบนโต๊ะ', indent=1.25)

# ── 2 · data table (template table 1) ───────────────────────────────────────
DATA_ROWS = [
    ['THEOS-2', 'GISTDA (AWAGAD)', 'Pan 0.5 ม. / MS 2 ม. Level 3 Ortho ก่อน–หลัง ก.ย. 67 รวมภาพ 17 ก.ย. 67',
     'ส่องรอยดินถล่มต้นน้ำ (L) · วัดความกว้างลำน้ำ (N) · ชี้เป้าจุดฟันหลอตลิ่ง 14 จุด · เทียบคราบโคลนรายหลังคาเรือนในองก์ 3'],
    ['Sphere API + ชั้นข้อมูลอาคาร BD_footprintTH', 'GISTDA',
     'Map API (thailand_images) · Routing API · อาคารทั้งประเทศ 24.5 ล้านหลัง ตัดเหลือแม่สาย 15,851 หลัง',
     'แผนที่ฐานดาวเทียมจริง · เส้นทางอพยพเลี่ยงถนนน้ำท่วมพร้อมจุดแวะอัตโนมัติ (เช่น สายลมจอย → ซอยเทศบาล 21 → ถ.พหลโยธิน 3.87 กม.) · คำนวณค่า P และผูกเคสเยียวยาจริงรหัส BD-<FID>'],
    ['เช็คน้ำ / Disaster Platform / LANDX', 'GISTDA',
     'เรดาร์ RADARSAT-2 วันที่ 10 และ 12 ก.ย. 67 · การใช้ประโยชน์ที่ดินต้นน้ำ',
     'ใช้ในองก์ 2 เมื่อเมฆบัง 100% · นับวันท่วมขัง ≥ 7 วัน (เกณฑ์เยียวยา) · LANDX ดูป่าต้นน้ำที่เปลี่ยนเป็นเหมือง/เกษตร'],
    ['ขอบเขต AOI · DEM · ถนน/แม่น้ำ', 'ทีม revenge · SRTM (NASA) · OpenStreetMap',
     'MaeSai_AOI_ALL.geojson (AOI-A/B) · SRTM DEM 30 ม. · โครงข่ายถนนและแม่น้ำสาย',
     'กรอบพื้นที่ภารกิจ · ความสูงพื้นดิน/มุมมอง 3D · โซนน้ำถึง · ตัวคำนวณเส้นทางสำรองแบบออฟไลน์'],
    ['ระดับน้ำ · ฝนต้นน้ำ · เกณฑ์เยียวยา', 'กรมทรัพยากรน้ำ · มูลนิธิอาสาเพื่อนพึ่ง (ภาฯ) ยามยาก ผ่าน ThaiWater · มติ ครม.',
     'สถานีบ้านถ้ำผาจม (วิกฤต 4.20 ม.) · ฝนสะสมสถานีบ้านโจตาดา เมียนมา · มติ ครม. 17 ก.ย. 67',
     'ค่า W และ Veto 1 · ค่า R ตัวชี้นำล่วงหน้า ~5 ชม. · อัตรา 9,000 + 10,000 บาทในองก์ 3'],
]
while len(dt.rows) < len(DATA_ROWS) + 1:
    dt._tbl.append(copy.deepcopy(dt.rows[-1]._tr))
for row, vals in zip(dt.rows[1:], DATA_ROWS):
    for cell, v in zip(row.cells, vals):
        fill_cell(cell, v, size=SMALL, align='left')
layout(dt, [3.0, 3.4, 4.9, 6.5])

# ── 3 · spatial analysis ─────────────────────────────────────────────────────
# the template forces a page break before section 3; with content filled in it leaves half a page empty
for br in find('3. Spatial Analysis')._p.iter(qn('w:br')):
    if br.get(qn('w:type')) == 'page':
        br.getparent().remove(br)
c = Cursor(find('(อธิบายขั้นตอนการนำข้อมูลมาพิจารณาร่วมกัน'))
c.table([
    ['INPUT', 'PROCESS', 'OUTPUT'],
    ['1. ภาพ THEOS-2 ก่อน–หลัง|2. ระดับน้ำถ้ำผาจม + ฝนโจตาดา|3. พื้นที่น้ำท่วมจากเรดาร์|4. อาคาร GISTDA 15,851 หลัง + DEM|5. ขอบเขต AOI + ถนน/แม่น้ำ',
     '1. นับรอยดินถล่ม วัดความกว้างลำน้ำ|2. ซ้อนทับอาคารกับระยะจากแม่น้ำ + DEM → โซนน้ำถึง|3. ปรับค่า W R L N P เป็น 0–1|4. คำนวณ SFVI ถ่วงน้ำหนัก + Veto|5. หลังน้ำลด: เทียบภาพรายหลังคา',
     '1. ลำดับชุมชนที่ต้องอพยพก่อน|2. เวลาที่เหลือก่อนน้ำถึงแต่ละชุมชน|3. เส้นทางอพยพเลี่ยงน้ำท่วม|4. บัญชีหลังคาเรือนที่ได้ 9,000 / +10,000 บาท|5. ตารางความผิดพลาดของการอนุมัติ'],
], [5.6, 6.3, 5.9])
c.para('**พิกเซลบริสุทธิ์ = (w − 2p) / w** (w = ความกว้างวัตถุ, p = ขนาดจุดภาพ) — รอยดินถล่ม 20 ม. ที่จุดภาพ 10 ม. เหลือพิกเซลบริสุทธิ์ 0% '
       'จึงตรวจพบได้แต่วัดความกว้างไม่ได้ จึงต้องใช้ THEOS-2', indent=1.25, keep=True)
c.table([
    ['ความกว้างรอย', 'จุดภาพ 10 ม. (Sentinel-2)', 'จุดภาพ 2 ม. (THEOS-2 MS)', 'จุดภาพ 0.5 ม. (THEOS-2 Pan)'],
    ['20 ม.', '0.0%', '80.0%', '95.0%'],
    ['25 ม.', '20.0%', '84.0%', '96.0%'],
    ['30 ม.', '33.3%', '86.7%', '96.7%'],
    ['จำนวนพิกเซล', '2–3', '10–15', '40–60'],
], [4.0, 4.6, 4.6, 4.6], align=['left', 'center', 'center', 'center'])
c.para('**การตัดข้อมูลอาคารจริง:** ตัดชั้น BD_footprintTH ทั้งประเทศด้วยกรอบ AOI-A ∩ อ.แม่สาย ได้ 15,851 หลัง คำนวณระยะแต่ละหลังจากแม่น้ำสายและความสูงจาก DEM '
       'หลังคาที่ห่างแม่น้ำ ≤ 150 ม. หรือ (≤ 800 ม. และพื้น ≤ 400 ม.รทก.) อยู่ใน “โซนน้ำถึง” รวม 3,071 หลัง แล้วจัดหลังคาให้ชุมชนที่ใกล้ที่สุดในรัศมี 300 ม. '
       '(เมื่อได้ขอบเขตน้ำท่วมเรดาร์จริงจะใช้แทนโซนจำลองทันที)', indent=1.25)

c = Cursor(find('3.2 ผลหรือหลักฐาน'))
c.para('ค่า P จริงจากอาคาร GISTDA เปลี่ยนลำดับความเสี่ยงจากที่ทีมตั้งไว้ในใบสมัคร (ค่าสถานการณ์) เป็นลำดับตามภูมิประเทศจริง:', indent=1.25, keep=True)
c.table([
    ['ชุมชน', 'หลังคาทั้งหมด', 'ในโซนน้ำถึง', 'ค่า P', 'ทางออก', 'ลำดับ SFVI'],
    ['เหมืองแดง', '340', '340', '1.00', '2', '1'],
    ['ไม้ลุงขน', '303', '295', '0.97', '3', '2'],
    ['เกาะทราย', '396', '339', '0.86', '2', '3'],
    ['สายลมจอย ม.1', '228', '120', '0.53', '1', '4'],
], [3.6, 2.8, 2.8, 2.4, 2.4, 3.8], align=['left', 'center', 'center', 'center', 'center', 'center'])
c.para('ผลที่ใช้ตัดสินใจ: (1) **เหมืองแดงและไม้ลุงขนเสี่ยงสูงสุด** เพราะหลังคาเกือบทั้งหมดอยู่ในโซนน้ำถึง จึงต้องได้รับทรัพยากรก่อน '
       '(2) สายลมจอยแม้ค่า P ต่ำสุด แต่เหลือทางออกเส้นเดียว จึงต้องมีเส้นทางสำรอง — Sphere Routing คำนวณได้ 3.87 กม. ผ่านซอยเทศบาล 21 สู่ ถ.พหลโยธิน '
       '(3) แนวตลิ่งฝั่งไทยยาว 2.8 กม. มีจุดฟันหลอ 14 จุด แต่ละจุดแสดงจำนวนหลังคาเสี่ยงด้านหลัง (≤ 150 ม.) เพื่อจัดลำดับการวางบิ๊กแบ็ก', indent=1.25)

c = Cursor(find('(อธิบายว่าข้อมูลใด สัมพันธ์'))
c.bullets([
    '**เหรียญข้อมูลจำกัด:** ชั้นข้อมูลทุกชั้นมีราคา ผู้เรียนต้องเลือกว่าจะ “ซื้อ” หลักฐานใดก่อน (เช่น ภาพ THEOS-2 + DEM + ชุมชน) จึงเห็นว่าการตัดสินใจที่ดีต้องใช้หลายชั้นร่วมกัน',
    '**ปิดชั้นภาพเชิงแสงอัตโนมัติ:** องก์ 2 เมฆ 100% ชั้น THEOS-2/LANDX ถูกล็อกต่อหน้า ผู้เรียนต้องหาหลักฐานทดแทนจากเรดาร์และ Sphere Routing',
    '**ป้ายที่มา 4 สี:** [วัดจริง] เขียว · [ภาพ THEOS-2] ฟ้า · [เรดาร์] ส้ม · [ค่าสถานการณ์] เทา กำกับทุกตัวเลขบนจอ เพื่อแยกข้อมูลจริงออกจากค่าจำลอง',
    '**ความสัมพันธ์ที่ต้องพิจารณาร่วม:** L (ดินถล่ม) บอกว่าน้ำจะพาโคลน · N (ลำน้ำแคบ) บอกว่าน้ำจะขึ้นเร็ว · P (หลังคาในโซนน้ำถึง) บอกว่าใครเสี่ยง — รวมกันเป็น SFVI รายชุมชน',
], size=BODY)

c = Cursor(find('ระบุว่าผู้เรียนต้องเลือกอะไร'))
c.para('**SFVI = 0.30·W + 0.25·R + 0.20·L + 0.15·N + 0.10·P**   (SFVI ≥ 0.62 = มีหลักฐานพอสั่งอพยพ)', indent=1.25, align='left', keep=True)
c.table([
    ['ตัวแปร', 'นิยาม', 'น้ำหนัก', 'ที่มา (ป้าย)'],
    ['W', 'ระดับน้ำสถานีบ้านถ้ำผาจม ÷ ระดับวิกฤต 4.20 ม.', '0.30', 'วัดจริง · กรมทรัพยากรน้ำ'],
    ['R', 'ฝนสะสมสถานีบ้านโจตาดา (เมียนมา) ÷ 250 มม.', '0.25', 'วัดจริง · ThaiWater'],
    ['L', 'รอยดินถล่มใหม่ที่ THEOS-2 พบ ÷ 17 จุด', '0.20', 'ภาพ THEOS-2'],
    ['N', '1 − ความกว้างลำน้ำที่วัดได้ ÷ 150 ม.', '0.15', 'ภาพ THEOS-2'],
    ['P', 'หลังคาจริงในโซนน้ำถึง ÷ หลังคาทั้งหมดของชุมชน', '0.10', 'อาคาร GISTDA + ค่าสถานการณ์'],
], [1.6, 9.4, 1.8, 5.0], align=['center', 'left', 'center', 'left'])
c.bullets([
    '**Veto 1:** ระดับน้ำถ้ำผาจม ≥ 4.20 ม. → สั่งอพยพทันทีไม่รอภาพ ปุ่ม “ยังไม่สั่งการ” ถูกปิด (4.01 ม. = เตรียมพร้อม)',
    '**Veto 2:** เมฆบัง > 60% หรือภาพเก่าเกิน 48 ชม. → ลดน้ำหนัก L จาก 0.20 เหลือ 0.08 และติดป้ายความเชื่อมั่นต่ำ',
    '**Trust Index:** เตือนผิดเสียความเชื่อมั่นแบบทบ −5 / −15 / −40 · เตือนช้าเสีย −30 ครั้งเดียวแต่มีผู้ได้รับผลกระทบจริง · คะแนนเท่ากันให้ชุมชนที่เหลือทางออกน้อยกว่าก่อน',
])
LOGIC = [
    ['ระดับน้ำ ≥ 4.20 ม.', 'สถานีถ้ำผาจม [วัดจริง]', 'ถึงระดับวิกฤต (Veto 1)', 'สั่งอพยพทันที ไม่รอภาพยืนยัน'],
    ['ระดับน้ำ 4.01–4.20 ม. และฝนต้นน้ำยังตก', 'W + R', 'มีแนวโน้มถึงวิกฤตในไม่กี่ชั่วโมง', 'แจ้งเตรียมพร้อม + สั่งถ่ายภาพรอบเร่งด่วน'],
    ['พบรอยดินถล่มใหม่ในพื้นที่รับน้ำ', 'THEOS-2 Pan 0.5 ม. (L)', 'น้ำจะพาตะกอนโคลนมาด้วย', 'เตรียมทรัพยากรล้างโคลน ไม่ใช่เฉพาะเรือ'],
    ['ลำน้ำเหลือ < ครึ่งของ 150 ม.', 'THEOS-2 วัดความกว้าง (N)', 'น้ำขึ้นเร็วกว่าเดิมที่ปริมาณเท่ากัน', 'ลดเกณฑ์เตือนภัย 0.20 ม. (4.01 → 3.81) เตือนเร็วขึ้น'],
    ['เมฆปกคลุม > 60% / ภาพเก่า > 48 ชม.', 'ภาพเชิงแสง + เมฆ', 'ค่า L ล้าสมัย (Veto 2)', 'ลดน้ำหนัก L · เมฆ 100% สลับใช้เรดาร์'],
    ['ชุมชนเหลือทางออกเส้นเดียว', 'ถนน + Sphere Routing', 'อพยพช้าแล้วออกไม่ได้', 'ยกลำดับขึ้น (tie-break) + คำนวณเส้นทางเบี่ยง'],
    ['ผู้เรียนเลือก “ยังไม่สั่งการ”', 'SFVI + เวลาน้ำถึง', 'รอข้อมูลแต่เวลาเดินต่อ 3 ชม.', 'แสดงผลเทียบสองทางของทั้ง 4 ชุมชน'],
    ['หลังน้ำลด: โคลนทับถมรอบอาคาร', 'THEOS-2 ก่อน–หลัง รายหลังคา', 'อยู่อาศัยไม่ได้จนกว่าจะล้างโคลน', 'อนุมัติค่าล้างโคลนเพิ่ม 10,000 บาท'],
    ['หลังน้ำลด: ท่วมขัง ≥ 7 วัน', 'ไทม์ซีรีส์เรดาร์', 'เข้าเกณฑ์ช่วยเหลือ', 'อนุมัติเหมาจ่าย 9,000 บาท'],
    ['หลังน้ำลด: หลักฐานไม่พอ', 'ภาพ + เรดาร์ (ความมั่นใจต่ำ)', 'ตัดสินจากหลักฐานที่มีไม่ได้', 'ส่งสำรวจภาคพื้น (แม่นแต่ช้า) / อนุมัติ (เสี่ยงจ่ายเกิน) / ปฏิเสธ (เสี่ยงตัดสิทธิ์ผิด)'],
]
while len(lm.rows) < len(LOGIC) + 1:
    lm._tbl.append(copy.deepcopy(lm.rows[-1]._tr))
for k, cell in enumerate(lm.rows[0].cells):
    fill_cell(cell, ['เกณฑ์/เงื่อนไข', 'ข้อมูลที่ใช้', 'ถ้าเกิดเงื่อนไขนี้', 'การตัดสินใจ'][k], size=SMALL, bold=True, white=True, align='center')
    shade(cell, NAVY)
for row, vals in zip(lm.rows[1:], LOGIC):
    for cell, v in zip(row.cells, vals):
        fill_cell(cell, v, size=SMALL)
layout(lm, [4.6, 3.8, 4.2, 5.2])

# ── 4 · mission design ──────────────────────────────────────────────────────
c = Cursor(find('4.1 เป้าหมายของภารกิจ'))
c.para('ให้ผู้เรียนวิเคราะห์หลักฐานหลายชั้น เปรียบเทียบทางเลือก และตัดสินใจสั่งการภายใต้เวลาและงบข้อมูลที่จำกัด พร้อมอธิบายเหตุผลด้วยค่า SFVI และเกณฑ์ยับยั้ง', indent=1.25)
c = Cursor(find('4.2 สิ่งที่ผู้เรียนจะได้รับ'))
c.bullets(['เข้าใจความจริงข้ามพรมแดน: ต้นเหตุของน้ำท่วมแม่สายอยู่นอกประเทศ และดาวเทียมคือเซนเซอร์เดียวที่มองเห็น',
           'เข้าใจข้อจำกัดของดาวเทียมเชิงแสงกับเรดาร์ SAR และเลือกใช้ให้ถูกช่วงเวลา',
           'เข้าใจต้นทุนของการรอข้อมูล (Cost of Inaction) และความเป็นธรรมของการเยียวยาด้วยหลักฐานดาวเทียม'])
c = Cursor(find('4.3 สถานการณ์/โจทย์เริ่มต้น'))
c.para('ผู้เรียนสวมบทศูนย์บัญชาการ EOC อำเภอแม่สาย ช่วงพายุไต้ฝุ่นยางิ ก.ย. 2567 เริ่มที่ T−48 ชม. ก่อนมวลน้ำถึงตัวเมือง ฝนตกหนักต่อเนื่องในต้นน้ำเมียนมาที่ไม่มีสถานีวัด '
       'มีโควตาสั่งถ่าย THEOS-2 เพียง 3 แถบ (แถบละ 10.3 กม.) แต้มสั่งการ 10,000 AP และเหรียญข้อมูล 14 เหรียญ', indent=1.25)
c = Cursor(find('4.4 คำถามเชิงพื้นที่หลัก'))
c.bullets(['**องก์ 1:** ควรสแกนป่าต้นน้ำ Sector ใดก่อน และควรอุดจุดฟันหลอของตลิ่งจุดใด เพราะหลังคาเสี่ยงด้านหลังมากที่สุด',
           '**องก์ 2:** เมื่อเมฆบัง 100% จะใช้หลักฐานอะไรแทน เส้นทางอพยพใดยังผ่านได้ และควรช่วยชุมชนใดก่อน',
           '**องก์ 3:** บ้านรหัส BD-<FID> หลังใดได้สิทธิ์ 9,000 บาท หลังใดควรได้เพิ่ม 10,000 บาท และหลังใดต้องส่งสำรวจ'])
c = Cursor(find('4.5 ลำดับการทำภารกิจ'))
flow = c.table([['Pre-test|5 ข้อ', 'องก์ 1 · T−48|เฝ้าระวัง: สแกนต้นน้ำ อุดตลิ่ง ตัดสินใจ', 'องก์ 2 · T+0|วิกฤต: เรดาร์ SAR + Sphere Routing + จัดลำดับ',
                 'องก์ 3 · T+36|ตรวจโคลนรายหลังคา + อนุมัติงบเยียวยา', 'AAR|+ Post-test']],
               [2.4, 4.4, 4.4, 4.2, 2.4], header=False, align=['center'] * 5)
for i, cell in enumerate(flow.rows[0].cells):
    shade(cell, ['E8EEF4', 'DDF5E6', 'FBE1E6', 'DCEEFB', 'E8EEF4'][i])
c.para('ลูกศรของภารกิจ: Pre-test → องก์ 1 → (หน้าสรุปผลการตัดสินใจ) → องก์ 2 → องก์ 3 → AAR & Post-test · ต้องตัดสินใจก่อนจึงผ่านไปองก์ถัดไปได้', indent=1.25, size=SMALL)

# ── 5 · prototype ───────────────────────────────────────────────────────────
c = Cursor(find('(อธิบายให้เข้าใจว่า Prototype'))
p512 = find('5.1 รูปแบบ Prototype')
# the template puts 5.1 and 5.2 in one paragraph (line break) — split it so each gets its own content
p52 = copy.deepcopy(p512._p)
kids = [k for k in p512._p if k.tag == qn('w:r')]
br_idx = next(i for i, k in enumerate(kids) if k.find(qn('w:br')) is not None)
for k in kids[br_idx:]: p512._p.remove(k)
kids2 = [k for k in p52 if k.tag == qn('w:r')]
for k in kids2[:br_idx + 1]: p52.remove(k)
c = Cursor(p512)
c.para('**Single-Page Interactive Geospatial Serious Game & Decision Simulator** — เว็บแอปไฟล์เดียว เปิดในเบราว์เซอร์ได้ทันที ไม่ต้องติดตั้งหรือล็อกอิน '
       'ทำงานแบบออนไลน์ (เชื่อม Sphere API: ภาพดาวเทียมและเส้นทาง) และออฟไลน์ (ข้อมูลที่ฝังในไฟล์ + ตัวคำนวณเส้นทางสำรอง)', indent=1.25)
c.put(p52)
c.bullets(['**เริ่ม:** อ่าน Mission Briefing → แบบทดสอบก่อนเรียน 5 ข้อ → เลือก **Beginner Mode** (ป้าย 2 บรรทัด: ภาษาไทยเข้าใจง่ายคู่ศัพท์เทคนิค พร้อมบอกว่าใช้ทำอะไร และมีแถบผู้ช่วยสอน) หรือ **Expert Mode** (ศัพท์เทคนิคล้วน ไม่มีตัวช่วย)',
           '**ทำ:** ทำตามกล่องภารกิจ 3 สเต็ปขององก์ 1 → องก์ 2 → องก์ 3 ใช้ **เหรียญข้อมูล** ซื้อชั้นข้อมูล และ **แต้มสั่งการ (AP)** สั่งถ่ายภาพ อุดตลิ่ง อพยพ คำนวณเส้นทาง',
           '**ได้:** ค่า SFVI เปลี่ยนสดทุกครั้งที่ได้หลักฐานใหม่ ทุกคำสั่งเดินนาฬิกาภารกิจ ผลการตัดสินใจสรุปเป็นคะแนนและหน้า AAR'])
c = Cursor(find('5.3 ส่วนประกอบหลักของ'))
c.table([
    ['ส่วนประกอบ', 'หน้าที่'],
    ['Top Bar', 'นาฬิกาภารกิจ (T−48 → T+36) คู่วันที่จริง ก.ย. 67 · หลังคาเรือนที่ยังเสี่ยง · แถบ 3 องก์ · ปุ่มไปองก์ถัดไป (ล็อกจนกว่าจะตัดสินใจ)'],
    ['Unified Data Hub (ซ้าย)', 'LIVE SFVI + ตัวแปร W R L N P แถวเดียวพร้อมป้ายที่มา · มาตรวัดระดับน้ำเทียบ 4.01/4.20 ม. · Trust Index'],
    ['Interactive Map (กลาง)', '3 มุมมอง: ลุ่มน้ำต้นน้ำ AOI-B (แถบถ่าย THEOS-2) · 4 ชุมชน AOI-A · แนวตลิ่งเมือง (ภาพ Sphere + อาคารจริง + 14 จุดฟันหลอ) และภูมิประเทศ 3D จาก DEM'],
    ['Single Action Toolbar (ขวา)', 'กล่องภารกิจ 3 สเต็ป (ติ๊กเขียวเมื่อเสร็จ) และแท็บ (1) ภาพต้นน้ำ (2) แนวตลิ่ง (3) ตัดสินใจ · องก์ 2: เส้นทาง/จัดลำดับ'],
    ['Interactive Tutor Bar (ล่าง)', 'บอกขั้นตอนถัดไปทีละข้อ พร้อมไฟชี้เป้า (Highlight) กระพริบที่ปุ่มที่ต้องกด'],
    ['AAR', 'คะแนนภารกิจ 5 ด้าน · ไทม์ไลน์การตัดสินใจ · บัญชีหลังคาเรือน · งบเยียวยา · บทเรียน 3 ข้อ · ส่งออกรายงาน CSV'],
], [4.2, 13.6], shade_first_col=True)

c = Cursor(find('5.4 รูปภาพ/หน้าจอ'))
W_IMG = 8.55
c.table([
    [('img', f'{IMG}/act1.jpg', W_IMG, '**ภาพที่ 1 องก์ 1** แนวตลิ่งเมือง: วัดลำน้ำ (N) + หมุดจุดฟันหลอพร้อมการ์ดสั่งอุด|ที่มา: GISTDA, ภาพดาวเทียม Sphere (thailand_images) และชั้นอาคาร BD_footprintTH; แม่น้ำ © OpenStreetMap'),
     ('img', f'{IMG}/act2.jpg', W_IMG, '**ภาพที่ 2 องก์ 2** เส้นทางอพยพสายลมจอย 3.87 กม. เลี่ยงถนนน้ำท่วม|ที่มา: GISTDA, Sphere Routing API และ Sphere Map API; โซนน้ำท่วมเป็นค่าสถานการณ์')],
    [('img', f'{IMG}/act3.jpg', W_IMG, '**ภาพที่ 3 องก์ 3** สไลเดอร์ก่อน/หลังรายหลังคา BD-<FID> + ปุ่มอนุมัติ 9,000 / +10,000 บาท|ที่มา: GISTDA, ภาพ Sphere (ก่อน) และ BD_footprintTH; ภาพหลังเป็นภาพจำลอง'),
     ('img', f'{IMG}/aar.jpg', W_IMG, '**ภาพที่ 4 AAR** คะแนนภารกิจ ไทม์ไลน์ และงบเยียวยา (Perfect Run 100%)|ที่มา: ผลจากต้นแบบของทีม revenge; ภาพก่อน–หลังเป็นภาพจำลองจนกว่าจะได้ THEOS-2 จริง')],
], [8.9, 8.9], header=False)
c.para('**ลิงก์เข้าใช้งาน Web Prototype:** ......................................................................  (ไฟล์ออฟไลน์: dist/index.html เปิดได้ทุกเบราว์เซอร์)', indent=1.25, align='left')

c = Cursor(find('5.5 ความเชื่อมโยงกับ Mission'))
c.table([
    ['ขั้นของ Mission', 'ส่วนของต้นแบบที่รองรับ'],
    ['รับโจทย์', 'Mission Briefing + กล่องภารกิจ 3 สเต็ปบอกเป้าหมายขององก์'],
    ['ใช้ข้อมูล', 'ชั้นข้อมูลแบบมีราคา (เหรียญ) · สั่งถ่าย THEOS-2 · ภาพ Sphere/อาคาร GISTDA · เรดาร์ในองก์ 2'],
    ['วิเคราะห์', 'LIVE SFVI รายชุมชน · ไม้บรรทัดวัดลำน้ำ · หลังคาเสี่ยงหลังจุดฟันหลอ · สไลเดอร์ก่อน/หลัง'],
    ['เปรียบเทียบ', 'ปุ่ม “ยังไม่สั่งการ” แสดงผลสองทาง · ลำดับ SFVI + tie-break · เส้นทาง Sphere เทียบเส้นทางสำรอง'],
    ['ตัดสินใจ', 'อพยพ / เตรียมพร้อม / รอ · จัดลำดับขบวนกู้ภัย · อนุมัติ/ปฏิเสธ/ส่งสำรวจรายหลังคา'],
    ['สรุปผล', 'หน้าสรุปองก์ 1 · AAR คะแนน 5 ด้าน · Post-test · Log CSV สำหรับครู'],
], [3.4, 14.4], shade_first_col=True)

# ── 6 · trials ──────────────────────────────────────────────────────────────
c = Cursor(find('6.1 ทดลองกับใคร'))
c.table([
    ['รอบ', 'กลุ่มทดลอง', 'จำนวน', 'เวลา', 'ตัวชี้วัดที่เก็บ'],
    ['1', 'เพื่อนร่วมชั้นที่ไม่เคยเรียน GIS', '8 คน', '45 นาที', 'เวลาสอนกติกา · ชุมชนที่เลือกอพยพเทียบกับ SFVI (ระบบบันทึกเอง)'],
    ['2', 'นักเรียน ม.ปลายที่ไม่เคยเรียน GIS', '10 คน', '60 นาที', 'Pre/Post-test 5 ข้อ · จำนวนชั้นข้อมูลที่เปิดก่อนตัดสินใจ'],
    ['3', 'ผู้ชมแหล่งเรียนรู้ (ทดสอบแบบบูธ)', '10 คน', '15 นาที/คน', 'เล่นจบภายในเวลาหรือไม่ · จุดที่สับสนหรือเลิกกลางคัน'],
], [1.2, 5.0, 1.6, 2.2, 7.8], align=['center', 'left', 'center', 'center', 'left'])
c.para('**เกณฑ์วัดผล:** สอนกติกา < 10 นาที · คะแนน Post-test เพิ่ม ≥ 30% · ผู้เล่น ≥ 70% เปิดข้อมูล ≥ 3 ชั้นก่อนตัดสินใจ (วัดจาก Log อัตโนมัติ)', indent=1.25)
c = Cursor(find('6.2 สิ่งที่เกิดขึ้นจากการทดลอง'))
c.para('**ทำได้ดี:** ผู้เรียนเข้าใจบทบาทของดาวเทียมในการมองข้ามพรมแดน และเหตุผลที่ต้องสลับไปใช้เรดาร์เมื่อเมฆบังได้ชัดเจน '
       '**จุดที่ต้องปรับ:** ในต้นแบบรุ่นแรก ผู้เรียนที่ไม่คุ้นเคยกับ GIS สับสนกับศัพท์เทคนิค หน้าจอมีแผงข้อมูลหลายส่วนจนไม่แน่ใจว่าต้องเริ่มกดปุ่มใดก่อน '
       'และใช้เวลานานในการค้นหาพิกัดบ้านในองก์ที่ 3', indent=1.25)
c = Cursor(find('6.3 สิ่งที่ทีมปรับหลังการทดลอง'))
c.table([
    ['ก่อน', 'หลัง', 'เหตุผล'],
    ['เมนูเป็นศัพท์เทคนิคชั้นเดียว', 'Beginner Mode: ป้าย 2 บรรทัด (ภาษาไทยเข้าใจง่าย + ศัพท์เทคนิคในวงเล็บ) พร้อมบอกว่าใช้ทำอะไร', 'ลดความสับสนของผู้ไม่เคยเรียน GIS'],
    ['ไม่มีลำดับนำทางชัดเจน', 'แถบผู้ช่วยสอน (Interactive Tutor Bar) ด้านล่าง + ไฟชี้เป้ากระพริบที่ขั้นตอนถัดไป + กล่องภารกิจ 3 สเต็ป', 'ผู้เรียนรู้ว่าต้องกดอะไรก่อน–หลังภายใน 30 วินาที'],
    ['ต้องเลื่อนแผนที่หาบ้านที่จะเยียวยาเอง', '1-Click Auto-Focus: คลิกรหัสบ้าน BD-<FID> แล้วซูมและเปิดสไลเดอร์ Before/After อัตโนมัติ', 'ลดเวลาในองก์ 3 ให้เวลาไปอยู่ที่การพิจารณาหลักฐาน'],
    ['มุมมองตลิ่งเป็นพื้นหลังจำลองและอาคารสมมติ', 'ภาพดาวเทียมจริงจาก Sphere API ซ้อนด้วยหลังคาเรือนจริงของ GISTDA 15,851 หลัง', 'ผู้เรียนเห็นแนวแม่น้ำและบ้านจริงใต้หมุดฟันหลอ'],
], [4.2, 8.4, 5.2])

# ── 7 · feasibility ─────────────────────────────────────────────────────────
c = Cursor(find('7.1 สถานที่ / บริบท'))
c.para('ห้องเรียน ม.ปลาย (วิชาภูมิศาสตร์/วิทยาการคำนวณ) · ศูนย์เรียนรู้และพิพิธภัณฑ์ เช่น Space Inspirium · ชุดนิทรรศการ GISTDA Roadshow · '
       'การซ้อมแผนบนโต๊ะ (Tabletop Exercise) ของหน่วยงานท้องถิ่น ขยายไปลุ่มน้ำอื่นได้โดยเปลี่ยนไฟล์ข้อมูลเชิงพื้นที่และค่าในไฟล์ตั้งค่า', indent=1.25)
c = Cursor(find('7.2 เวลา อุปกรณ์'))
c.para('**เวลา** 15 นาที (บูธ) ถึง 45 นาที (ห้องเรียน) ต่อรอบ · **อุปกรณ์** คอมพิวเตอร์หรือแท็บเล็ตทั่วไปที่เปิดเว็บเบราว์เซอร์ได้ รองรับ Offline Mode ไม่ต้องพึ่งอินเทอร์เน็ตหน้างาน '
       '· **พื้นที่** โต๊ะ 1 ตัวต่อผู้เล่น 1–2 คน · **บุคลากร** ผู้ดูแล 1 คน (ส่งออก Log ได้จากโหมดครู)', indent=1.25)

# ── 8 · creativity ──────────────────────────────────────────────────────────
c = Cursor(find('8.1 จุดแตกต่าง'))
c.para('ผสานข้อมูลจริงหลายเซนเซอร์ (THEOS-2 · เรดาร์ · สถานีวัด · อาคาร GISTDA · Sphere API) เข้ากับกลไกความกดดันด้านเวลา ต้นทุนข้อมูล และความเชื่อมั่นประชาชน (Trust Index) '
       'จำลองภาวะเมฆบัง 100% ที่บังคับให้เปลี่ยนเซนเซอร์ต่อหน้า และให้เยียวยาบนพิกัดอาคารจริงรายหลัง แทนตัวเลขสมมติ', indent=1.25)
c = Cursor(find('8.2 จุดเด่นของ Prototype'))
c.bullets(['**ป้ายที่มา 4 สี** สร้างความฉลาดรู้ด้านข้อมูล (Data Literacy) แยกข้อมูลวัดจริง ภาพดาวเทียม เรดาร์ และค่าสถานการณ์ออกจากกัน',
           '**ปุ่ม “ยังไม่สั่งการ”** แสดงผลเปรียบเทียบสองทางของทั้ง 4 ชุมชน ให้เห็นราคาของการชะลอเวลาเป็นตัวเลข',
           '**หน้า AAR** วิเคราะห์เวลาที่ลังเลในแต่ละการตัดสินใจ พร้อมโหมดฝึกซ้ำเฉพาะจุดที่ทำผิด'])

# ── 9 · demo video ──────────────────────────────────────────────────────────
c = Cursor(find('วางลิงก์ / QR Code'))
box = c.table([['**Mission Demo Video:** ........................................................................................|'
                '**Web Prototype:** ........................................................................................|'
                'วิดีโอสาธิตแบบ Perfect Run ครบ 3 องก์ พร้อมคำบรรยายภาษาไทยทุกขั้น (ความยาวประมาณ 5 นาที 40 วินาที)']],
              [17.8], header=False, size=BODY)

# drop the form's own submission instructions after section 9
tail = find('ผู้เข้าแข่งขันเสนอแนวคิดภารกิจ')
el = tail._p
while el is not None and el.tag != qn('w:sectPr'):
    nxt = el.getnext()
    el.getparent().remove(el)
    el = nxt
# trim blank paragraphs between the link box and the end
body = doc.element.body
last = body[-2] if body[-1].tag == qn('w:sectPr') else body[-1]
while last.tag == qn('w:p') and not ''.join(last.itertext()).strip() and last.getprevious() is not None and last.getprevious().tag == qn('w:p'):
    prev = last.getprevious(); body.remove(last); last = prev

doc.save(OUT)
print('saved', OUT)
